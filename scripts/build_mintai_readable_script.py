from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "script_1786448782605360400.json"
PROFILES = ROOT / "outputs" / "actors_profile_1786448782605360400.json"
OUT_DIR = ROOT / "outputs" / "readable"
DOCX_OUT = OUT_DIR / "共煮海峡一碗羹_可读分镜剧本.docx"
MD_OUT = OUT_DIR / "共煮海峡一碗羹_可读分镜剧本.md"
SIMPLE_MD_OUT = OUT_DIR / "共煮海峡一碗羹_剧本.md"

TITLE = "共煮海峡一碗羹"
SUBTITLE = "闽台文化温情短片 · 可读分镜剧本"
META = "两幕 · 莲花古镇 · 约 2 分钟"


def set_run_font(run, size=11, bold=None, italic=None, color=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run, size=9, color="777777")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("Scene Metadata", "Shot Direction", "Dialogue"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    meta = doc.styles["Scene Metadata"]
    meta.base_style = doc.styles["Normal"]
    meta.paragraph_format.space_after = Pt(5)
    meta.paragraph_format.keep_with_next = True

    direction = doc.styles["Shot Direction"]
    direction.base_style = doc.styles["Normal"]
    direction.paragraph_format.left_indent = Inches(0.22)
    direction.paragraph_format.right_indent = Inches(0.22)
    direction.paragraph_format.space_after = Pt(6)

    dialogue = doc.styles["Dialogue"]
    dialogue.base_style = doc.styles["Normal"]
    dialogue.paragraph_format.left_indent = Inches(0.35)
    dialogue.paragraph_format.right_indent = Inches(0.35)
    dialogue.paragraph_format.space_after = Pt(10)
    dialogue.paragraph_format.keep_together = True


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(112)
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    set_run_font(run, size=30, bold=True, color="203748")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(SUBTITLE)
    set_run_font(run, size=15, color="2B5163")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(86)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(META)
    set_run_font(run, size=10.5, italic=True, color="7A5A00")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("由 ScriptsGenerateAgent 多 Agent 流程生成")
    set_run_font(run, size=9.5, color="666666")
    doc.add_page_break()


def add_roles(doc, profiles):
    doc.add_heading("角色", level=1)
    backgrounds = {
        "秀枝": "68岁。泉州家庭出身，年轻时随家人迁居台湾；克制、务实，熟悉家族办桌做法。",
        "雨晴": "25岁。在台南长大，从事短视频美食记录；敏锐、好奇，尊重长辈但有主见。",
        "海生": "31岁。厦门人，带着祖父保存的旧菜谱来到古镇；认真温和，对家族记忆有执念。",
    }
    for profile in profiles:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        name = profile.get("name", "")
        r = p.add_run(f"{name}　")
        set_run_font(r, size=11, bold=True, color="1F4D78")
        r = p.add_run(backgrounds.get(name, profile.get("background", "")))
        set_run_font(r, size=11)


def event_kind(event):
    if not event.get("speaker") and not event.get("content"):
        return "空镜"
    if event.get("shot") == "move":
        return "调度"
    return "对白"


def action_text(event):
    pieces = []
    for action in event.get("actions") or []:
        detail = action.get("motion_detail") or action.get("action") or ""
        character = action.get("character") or ""
        if detail:
            pieces.append(f"{character}：{detail}" if character else detail)
    return "；".join(pieces)


def add_event(doc, event, index):
    kind = event_kind(event)
    shot_type = event.get("shot_type") or ("全景" if kind == "空镜" else event.get("shot", "镜头"))
    duration = event.get("duration") or "未标注"
    p = doc.add_paragraph(style="Scene Metadata")
    r = p.add_run(f"镜头 {index:02d}　{kind} · {shot_type} · {duration}")
    set_run_font(r, size=9.5, bold=True, color="7A5A00")

    description = event.get("shot_description") or ""
    action = action_text(event)
    if description or action:
        p = doc.add_paragraph(style="Shot Direction")
        text = description
        if action:
            text = f"{text} 动作：{action}" if text else f"动作：{action}"
        r = p.add_run(text)
        set_run_font(r, size=10.5, italic=True, color="555555")

    if event.get("speaker") or event.get("content"):
        p = doc.add_paragraph(style="Dialogue")
        speaker = event.get("speaker") or "旁白"
        r = p.add_run(f"{speaker}：")
        set_run_font(r, size=11, bold=True, color="203748")
        r = p.add_run(event.get("content") or "")
        set_run_font(r, size=11)


def build_docx(script, profiles):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = header.add_run("共煮海峡一碗羹 · 可读分镜剧本")
    set_run_font(r, size=9, color="777777")
    add_page_number(section.footer.paragraphs[0])

    add_cover(doc)
    add_roles(doc, profiles)

    for act_index, act in enumerate(script, 1):
        doc.add_heading(f"第 {act_index} 幕", level=1)
        info = act.get("scene information") or {}
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(f"场景：{info.get('where', '莲花古镇')}\n")
        set_run_font(r, size=10.5, bold=True, color="1F4D78")
        r = p.add_run(info.get("what", ""))
        set_run_font(r, size=10.5, italic=True, color="555555")

        for event_index, event in enumerate(act.get("scene") or [], 1):
            add_event(doc, event, event_index)

    core = doc.core_properties
    core.title = TITLE
    core.subject = "闽台文化温情短片可读分镜剧本"
    core.author = "ScriptsGenerateAgent"
    doc.save(DOCX_OUT)


def md_escape(text):
    return str(text or "").replace("|", "\\|")


def build_markdown(script, profiles):
    lines = [f"# {TITLE}", "", f"> {SUBTITLE}  ", f"> {META}", "", "## 角色", ""]
    backgrounds = {
        "秀枝": "68岁。泉州家庭出身，年轻时随家人迁居台湾；克制、务实，熟悉家族办桌做法。",
        "雨晴": "25岁。在台南长大，从事短视频美食记录；敏锐、好奇，尊重长辈但有主见。",
        "海生": "31岁。厦门人，带着祖父保存的旧菜谱来到古镇；认真温和，对家族记忆有执念。",
    }
    for profile in profiles:
        name = profile.get("name", "")
        lines.extend([f"- **{name}**：{backgrounds.get(name, profile.get('background', ''))}", ""])

    for act_index, act in enumerate(script, 1):
        info = act.get("scene information") or {}
        lines.extend([
            f"## 第 {act_index} 幕",
            "",
            f"**场景：** {info.get('where', '莲花古镇')}",
            "",
            f"*{info.get('what', '')}*",
            "",
        ])
        for event_index, event in enumerate(act.get("scene") or [], 1):
            kind = event_kind(event)
            shot_type = event.get("shot_type") or ("全景" if kind == "空镜" else event.get("shot", "镜头"))
            duration = event.get("duration") or "未标注"
            lines.extend([f"### 镜头 {event_index:02d}", "", f"`{kind}` `{shot_type}` `{duration}`", ""])
            description = event.get("shot_description") or ""
            action = action_text(event)
            if description:
                lines.extend([f"> {md_escape(description)}", ""])
            if action:
                lines.extend([f"> 动作：{md_escape(action)}", ""])
            if event.get("speaker") or event.get("content"):
                lines.extend([f"**{event.get('speaker') or '旁白'}：** {event.get('content') or ''}", ""])

    MD_OUT.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def build_simple_markdown(script):
    lines = [f"# {TITLE}", ""]
    chinese_act_names = ["第一幕", "第二幕", "第三幕", "第四幕", "第五幕"]

    for act_index, act in enumerate(script):
        act_name = chinese_act_names[act_index] if act_index < len(chinese_act_names) else f"第 {act_index + 1} 幕"
        info = act.get("scene information") or {}
        lines.extend([f"## {act_name}", "", f"*{info.get('where', '莲花古镇')}。{info.get('what', '')}*", ""])

        for event in act.get("scene") or []:
            description = (event.get("shot_description") or "").strip()
            if description:
                lines.extend([f"*{description}*", ""])

            if event.get("speaker") or event.get("content"):
                speaker = event.get("speaker") or "旁白"
                lines.extend([f"**{speaker}：** {event.get('content') or ''}", ""])

    SIMPLE_MD_OUT.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    script = json.loads(SOURCE.read_text(encoding="utf-8"))
    profiles = json.loads(PROFILES.read_text(encoding="utf-8"))
    build_markdown(script, profiles)
    build_simple_markdown(script)
    build_docx(script, profiles)
    print(DOCX_OUT)
    print(MD_OUT)
    print(SIMPLE_MD_OUT)


if __name__ == "__main__":
    main()
